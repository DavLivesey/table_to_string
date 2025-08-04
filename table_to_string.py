import openpyxl
import docx

path = '/path/to/get/file.xlsx'
exel = openpyxl.load_workbook(path)
doc = docx.Document('/path/to/edit/file.docx')

workers = exel['Sheet1']


def table_to_string():
    email_string = []
    for number_row in range(1, len(workers['A'])):
            email_string.append(str(workers.cell(row=number_row, column=1).value))
    doc.paragraphs[0].text = (",").join(email_string).split(", ")
    doc.save('/path/to/save/file.docx')


if __name__ == "__main__":
    table_to_string()
